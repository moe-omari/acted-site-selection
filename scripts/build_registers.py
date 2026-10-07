"""Bake verification assessments and ACT sites for area export."""

import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"d:\new - site verification assessments - act sites - managed sites.xlsx")
OUT = ROOT / "src" / "data" / "registers.json"


def json_cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="minutes")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return int(value)
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def sheet_table(workbook, name):
    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    header = ["" if cell is None else str(cell) for cell in next(rows)]
    body = []
    for row in rows:
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        values = [json_cell(cell) for cell in row]
        if len(values) < len(header):
            values.extend([None] * (len(header) - len(values)))
        body.append(values[: len(header)])
    return header, body


def index_of(header, label):
    try:
        return header.index(label)
    except ValueError as error:
        raise SystemExit(f"missing column {label}") from error


def main():
    workbook = load_workbook(XLSX, read_only=True, data_only=True)
    verification_header, verification_rows = sheet_table(workbook, "Sheet1")
    act_header, act_rows = sheet_table(workbook, "ACT sites")
    payload = {
        "verification": {
            "headers": verification_header,
            "rows": verification_rows,
            "siteId": index_of(verification_header, "Site ID"),
            "formSiteId": index_of(verification_header, "Form Site ID"),
        },
        "act": {
            "headers": act_header,
            "rows": act_rows,
            "siteId": index_of(act_header, "Site ID"),
            "newSiteId": index_of(act_header, "New Site ID"),
            "lon": index_of(act_header, "Longitude"),
            "lat": index_of(act_header, "Latitude"),
        },
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"verification {len(verification_rows)} rows, act {len(act_rows)} rows, bytes {OUT.stat().st_size}")


if __name__ == "__main__":
    main()
