"""Build src/data/partner-sites.json from the partner managed-sites workbook."""
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(r'C:\Users\PC\Desktop\managed sites by partners.xlsx')
TARGET = Path('src/data/partner-sites.json')

HOUSEHOLDS = 'Please review the below information / Estimated number of households currently accommodated in the site (total population divided by 5 is ${est_hh_existing})'
INDIVIDUALS = 'Estimated number of individuals currently accommodated in the site'


def number(value):
    if value is None or value == '':
        return None
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return None
    if amount != amount:
        return None
    if amount.is_integer():
        return int(amount)
    return round(amount, 1)


def coordinate(value):
    if value is None or value == '':
        return None
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return None
    if amount != amount:
        return None
    return amount


def text(value):
    return str(value or '').strip()


def main():
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(cell) for cell in next(rows)]
    index = {name: position for position, name in enumerate(headers)}
    seen = Counter()
    sites = []
    for row in rows:
        site_id = text(row[index['Site ID']])
        lat = coordinate(row[index['Latitude']])
        lon = coordinate(row[index['Longitude']])
        if not site_id or lat is None or lon is None:
            continue
        seen[site_id] += 1
        key = site_id if seen[site_id] == 1 else f'{site_id}-{seen[site_id]}'
        sites.append({
            'key': key,
            'id': site_id,
            'name': text(row[index['Site Name']]),
            'nameAr': text(row[index['Site Name (Arabic)']]),
            'partner': text(row[index['Implementing Partner']]),
            'governorate': text(row[index['First Level Region Name']]),
            'neighborhood': text(row[index['Second Level Region Name']]),
            'siteType': text(row[index['Site Type']]),
            'status': text(row[index['Site Status']]),
            'households': number(row[index[HOUSEHOLDS]]),
            'individuals': number(row[index[INDIVIDUALS]]),
            'lat': lat,
            'lon': lon,
        })
    TARGET.write_text(json.dumps(sites, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'{len(sites)} sites')


if __name__ == '__main__':
    main()
