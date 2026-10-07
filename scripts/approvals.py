"""Map approved ACT codes to the official site ID from the verification workbook."""

from pathlib import Path

from openpyxl import load_workbook

SITES_XLSX = Path(r"d:\new - site verification assessments - act sites - managed sites.xlsx")


def approved_site_codes(path=SITES_XLSX):
    """Form Site ID (ACT code) -> Site ID, for active verification assessments."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell) if cell is not None else "" for cell in next(rows)]
    site_index = header.index("Site ID")
    form_index = header.index("Form Site ID")
    status_index = header.index("Site Status")
    alias = {}
    for row in rows:
        if not row or not row[form_index] or not row[site_index]:
            continue
        status = str(row[status_index] or "")
        if not status.startswith("Active"):
            continue
        form_id = str(row[form_index]).strip().upper()
        site_id = str(row[site_index]).strip().upper()
        if not form_id.startswith("ACT") or form_id == site_id:
            continue
        alias[form_id] = site_id
    return alias
