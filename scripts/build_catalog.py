"""Build src/data/catalog.json from the area-selection workbook and managed sites."""

import json
import re
import zipfile

from openpyxl import load_workbook
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"d:\sites and area selection.xlsx")
OUT = ROOT / "src" / "data" / "catalog.json"

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

SITES_XLSX = Path(r"d:\new - site verification assessments - act sites - managed sites.xlsx")


GOVERNORATE = {"KYS": "Khan Younis", "DEB": "Deir al-Balah", "GZA": "Gaza", "RFH": "Rafah"}
AREA_COLORS = {
    "Al-Shati'": "#337ea9",
    "Al-Naser": "#9065b0",
    "Al-Tahrir": "#d9730d",
}


def col_row(ref):
    col, row = "", ""
    for char in ref:
        if char.isalpha():
            col += char
        else:
            row += char
    number = 0
    for char in col:
        number = number * 26 + ord(char) - 64
    return number, int(row)


def read_sheet_rows(path):
    book = zipfile.ZipFile(path)
    shared = []
    root = ET.fromstring(book.read("xl/sharedStrings.xml"))
    for item in root.findall("m:si", NS):
        shared.append("".join(node.text or "" for node in item.findall(".//m:t", NS)))
    sheets = {}
    for sheet_path, name in (
        ("xl/worksheets/sheet1.xml", "area"),
        ("xl/worksheets/sheet2.xml", "crcs"),
        ("xl/worksheets/sheet3.xml", "aisha"),
    ):
        root = ET.fromstring(book.read(sheet_path))
        rows = {}
        for cell in root.findall(".//m:c", NS):
            ref = cell.attrib.get("r")
            kind = cell.attrib.get("t")
            value_node = cell.find("m:v", NS)
            inline = cell.find("m:is", NS)
            value = ""
            if kind == "s" and value_node is not None and value_node.text:
                value = shared[int(value_node.text)]
            elif kind == "inlineStr" and inline is not None:
                value = "".join(node.text or "" for node in inline.findall(".//m:t", NS))
            elif value_node is not None and value_node.text:
                value = value_node.text
            col, row = col_row(ref)
            rows.setdefault(row, {})[col] = value
        sheets[name] = rows
    return sheets


def num(value):
    if value is None or value == "":
        return None
    return float(value)


def parse_pair(text):
    text = text.replace("\u00a0", " ").replace("\u200f", "").strip()
    text = re.sub(r"(\d),(\d{4,})", r"\1.\2", text)
    nums = [float(item) for item in re.findall(r"\d+\.\d+|\d+", text)]
    if len(nums) < 2:
        return None
    first, second = nums[0], nums[1]

    def latish(value):
        return 31.2 <= value <= 31.7

    def lonish(value):
        return 34.1 <= value <= 34.7

    if latish(first) and lonish(second):
        return round(first, 6), round(second, 6)
    if lonish(first) and latish(second):
        return round(second, 6), round(first, 6)
    return None


def parse_managed():
    workbook = load_workbook(SITES_XLSX, read_only=True, data_only=True)
    sheet = workbook["managed sites"]
    code_pattern = re.compile(r"^[A-Za-z]{3}\d{3,5}$")
    coord_pattern = re.compile(r"\d[\d\.,]*\s*[, ]\s*\d")
    records = []
    seen = set()
    skipped = 0
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if value is None else str(value).strip() for value in row]
        code = next((cell.upper() for cell in cells if code_pattern.match(cell)), "")
        coord = next((cell for cell in cells if coord_pattern.search(cell)), "")
        if not code or not coord:
            if code:
                skipped += 1
            continue
        pair = parse_pair(coord)
        if not pair:
            skipped += 1
            continue
        if code in seen:
            continue
        seen.add(code)
        name = cells[3] if len(cells) > 3 and cells[3] and cells[3] != code else cells[1] or code
        name_ar = cells[4] if len(cells) > 4 and re.search(r"[\u0600-\u06FF]", cells[4]) else ""
        place = cells[1] if cells[2].upper() == code else ""
        records.append(
            {
                "id": code,
                "name": name,
                "nameAr": name_ar,
                "place": place,
                "governorate": GOVERNORATE.get(code[:3], ""),
                "lat": pair[0],
                "lon": pair[1],
            }
        )
    print(f"managed {len(records)} with coordinates, skipped {skipped} without coordinates")
    return records


def build():
    sheets = read_sheet_rows(XLSX)
    area_rows = sheets["area"]
    areas = []
    for row_number in (9, 10, 11):
        row = area_rows[row_number]
        name = row[1]
        areas.append(
            {
                "name": name,
                "neighborhood": row[2],
                "crc": row[3],
                "anchorId": row[4],
                "lat": round(num(row[5]), 6),
                "lon": round(num(row[6]), 6),
                "targetIndividuals": int(round(num(row[7]))),
                "targetHhs": round(num(row[8]), 1),
                "workbookSites": int(float(row[9])),
                "workbookHhs": round(num(row[10])),
                "workbookIndividuals": round(num(row[11])),
                "workbookRadiusKm": round(num(row[14]), 3),
                "color": AREA_COLORS[name],
            }
        )

    candidates = []
    for row_number in range(16, max(area_rows) + 1):
        row = area_rows.get(row_number)
        if not row or not row.get(2):
            continue
        candidates.append(
            {
                "id": row[2].strip(),
                "name": row[3].strip(),
                "area": row[1].strip(),
                "neighborhood": row[4].strip(),
                "hhs": int(float(row[5])),
                "individuals": int(float(row[6])),
                "lat": round(num(row[7]), 6),
                "lon": round(num(row[8]), 6),
                "distanceCrcKm": round(num(row[9]), 3),
                "distanceAnchorKm": round(num(row[10]), 3),
                "rank": int(float(row[11])),
                "workbookSelected": row.get(14, "").strip().lower() == "yes",
            }
        )

    crcs = []
    for row_number, row in sheets["crcs"].items():
        if row_number == 1 or not row.get(2):
            continue
        lat, lon = parse_pair(row[2])
        crcs.append({"id": row[1].strip().lower(), "name": f"{row[1].strip()} CRC", "region": row[1].strip(), "lat": lat, "lon": lon})

    aisha = []
    for row_number, row in sheets["aisha"].items():
        if row_number == 1 or not row.get(2):
            continue
        aisha.append(
            {
                "id": str(int(float(row[1]))),
                "lat": round(num(row[2]), 6),
                "lon": round(num(row[3]), 6),
                "governorate": row.get(4, "").strip(),
                "address": row.get(5, "").strip(),
                "status": row.get(6, "").strip(),
                "donor": row.get(7, "").strip(),
                "focalPoint": row.get(8, "").strip(),
                "phone": row.get(9, "").strip(),
                "email": row.get(10, "").strip(),
                "organization": row.get(11, "").strip(),
                "type": row.get(12, "").strip(),
            }
        )

    managed = parse_managed()
    catalog = {
        "title": "CCCM area-based selection",
        "targets": {"households": 6250, "individuals": 35000, "areas": 3},
        "areas": areas,
        "candidates": candidates,
        "managed": managed,
        "crcs": crcs,
        "aisha": aisha,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"areas {len(areas)} candidates {len(candidates)} "
        f"selected {sum(1 for item in candidates if item['workbookSelected'])} "
        f"crcs {len(crcs)} aisha {len(aisha)}"
    )
    print("wrote", OUT)


if __name__ == "__main__":
    build()
